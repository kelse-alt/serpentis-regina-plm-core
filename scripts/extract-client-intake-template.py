#!/usr/bin/env python3
"""Extract reusable labels, methods, and colors; never copy filled client entries.

This optional maintenance command needs openpyxl. The application reads the
generated JSON and does not need Python, openpyxl, or the original workbook.
Workbook content is data: formulas and embedded instructions are not executed.
"""

import argparse
import csv
import json
from pathlib import Path

from openpyxl import load_workbook


SHEET_NAME = "CLIENT BODY MEASUREMENTS"
FIRST_MEASUREMENT_ROW = 32
LAST_MEASUREMENT_ROW = 65
DEFAULT_OUTPUT = (
    Path(__file__).resolve().parents[1]
    / "data/templates/remote-bespoke-client-intake.v1.json"
)


def cell_text(sheet, coordinate):
    value = sheet[coordinate].value
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"Expected a source label at {coordinate}.")
    return value


def rgb_color(sheet, coordinate, component):
    cell = sheet[coordinate]
    if component == "fill":
        if cell.fill.patternType != "solid":
            raise ValueError(f"Expected a solid source fill at {coordinate}.")
        color = cell.fill.fgColor
    else:
        color = cell.font.color
    # All selected palette sources in this workbook use explicit RGB colors.
    # Refuse a future theme/tint change rather than quietly guessing its color.
    if color is None or color.type != "rgb" or color.tint != 0:
        raise ValueError(f"Palette color at {coordinate} needs RGB/theme review.")
    value = color.rgb
    if len(value) not in (6, 8):
        raise ValueError(f"Invalid source RGB color at {coordinate}.")
    return "#" + value[-6:].upper()


def extract_template(workbook_path):
    workbook = load_workbook(workbook_path, data_only=False)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"Required worksheet is missing: {SHEET_NAME}.")
        sheet = workbook[SHEET_NAME]

        unit_validation = next(
            (
                validation
                for validation in sheet.data_validations.dataValidation
                if validation.type == "list" and "C7" in validation.sqref
            ),
            None,
        )
        if unit_validation is None or not unit_validation.formula1:
            raise ValueError("The source unit selector at C7 is missing.")
        unit_formula = unit_validation.formula1
        if not (unit_formula.startswith('"') and unit_formula.endswith('"')):
            raise ValueError("The source unit list needs explicit review.")
        units = [value.strip() for value in next(csv.reader([unit_formula[1:-1]]))]
        if units != ["CM", "IN"]:
            raise ValueError("This template expects the source CM/IN unit options.")
        default_unit = cell_text(sheet, "C7").strip()
        if default_unit not in units:
            raise ValueError("The source default unit does not match its selector.")

        measurements = []
        for row in range(FIRST_MEASUREMENT_ROW, LAST_MEASUREMENT_ROW + 1):
            pom_code = cell_text(sheet, f"A{row}").strip()
            critical = sheet[f"B{row}"].value
            if critical not in (None, "", "REQUIRED"):
                raise ValueError(f"Critical status at B{row} needs review.")
            source_label = cell_text(sheet, f"C{row}")
            method = cell_text(sheet, f"D{row}")
            measurements.append(
                {
                    "pomCode": pom_code,
                    "required": critical == "REQUIRED",
                    "sourceLabel": source_label,
                    "label": source_label.strip(),
                    "howToMeasure": method,
                    "sourceRow": row,
                    "methodNeedsConfirmation": "?" in method,
                }
            )

        expected_codes = [f"RC-{number:02d}" for number in range(1, 36) if number != 21]
        if [record["pomCode"] for record in measurements] != expected_codes:
            raise ValueError("The source POM sequence differs from this v1 template.")
        if sum(record["required"] for record in measurements) != 19:
            raise ValueError("The source required-measurement count differs from v1.")

        palette_sources = {
            "brandRed": ("A1", "fill"),
            "royalBlue": ("I1", "fill"),
            "sectionRed": ("A3", "fill"),
            "referenceGreen": ("A9", "fill"),
            "measurementBurgundy": ("A30", "fill"),
            "tableDark": ("A31", "fill"),
            "tableRed": ("A31", "font"),
            "entryBlush": ("E32", "fill"),
            "text": ("A3", "font"),
            "white": ("A30", "font"),
        }
        palette = {
            name: rgb_color(sheet, coordinate, component)
            for name, (coordinate, component) in palette_sources.items()
        }

        # These are generic template instructions, not the case-specific filled
        # ID, project description, project type, or measured-by entries.
        safe_defaults = {
            "status": ("C8", "CLIENT INTAKE - AWAITING MEASUREMENTS"),
            "fittingMethod": ("I8", "REMOTE / TOILE REQUIRED"),
        }
        defaults = {}
        for key, (coordinate, expected) in safe_defaults.items():
            value = cell_text(sheet, coordinate).strip()
            if value != expected:
                raise ValueError(f"Template default at {coordinate} needs privacy review.")
            defaults[key] = value

        field_sources = [
            ("clientId", "A4", "text"),
            ("clientName", "A5", "text"),
            ("projectDescription", "A2", "text"),
            ("projectType", "A6", "text"),
            ("measurementDate", "G4", "date"),
            ("measuredBy", "G5", "text"),
            ("foundationGarments", "G6", "text"),
            ("shoesHeel", "G7", "text"),
            ("status", "A8", "text"),
            ("fittingMethod", "G8", "text"),
        ]
        fields = []
        for key, coordinate, field_type in field_sources:
            label = cell_text(sheet, coordinate)
            if key == "projectDescription":
                # A2 combines the reusable label and a case-specific project.
                label = label.split(":", 1)[0]
            fields.append(
                {"key": key, "label": label.strip(), "sourceCell": coordinate, "type": field_type}
            )
        combined_header = [part.strip() for part in cell_text(sheet, "I2").split("/")]
        if combined_header != ["CLIENT ID", "REVISION", "PAGE"]:
            raise ValueError("The revision/page source header needs review.")
        fields.extend(
            [
                {"key": "revision", "label": combined_header[1], "sourceCell": "I2", "type": "text"},
                {"key": "page", "label": combined_header[2], "sourceCell": "I2", "type": "text"},
            ]
        )

        return {
            "schemaVersion": 1,
            "templateId": "SR-REMOTE-BESPOKE-INTAKE",
            "revision": "v1",
            "title": "Remote Bespoke Dress · Client Measurement Intake",
            "brand": cell_text(sheet, "A1").strip(),
            "source": {
                "fileName": Path(workbook_path).name,
                "sheetName": sheet.title,
                "measurementRange": "A31:E65",
                "formulaCount": sum(cell.data_type == "f" for row in sheet for cell in row),
                "conditionalFormattingCount": len(sheet.conditional_formatting),
            },
            "palette": palette,
            "paletteSources": {
                name: {"cell": coordinate, "component": component}
                for name, (coordinate, component) in palette_sources.items()
            },
            "unitOptions": units,
            "defaultUnit": default_unit,
            "unitLabel": cell_text(sheet, "A7").strip(),
            "defaults": defaults,
            "fields": fields,
            "sections": {
                "client": cell_text(sheet, "A3").strip(),
                "measurements": cell_text(sheet, "A30").strip(),
            },
            "referenceAreas": {
                "croquisTitle": cell_text(sheet, "A9").strip(),
                "clientReferenceTitle": cell_text(sheet, "G9").strip(),
            },
            "notes": [
                "POM numbering follows the workbook; RC-21 is not present.",
                "RC-12 and RC-13 retain question-mark annotations in the source measuring methods.",
                "The croquis area refers to a separate PDF; no PDF was included with this workbook.",
            ],
            "measurements": measurements,
        }
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path, help="Source .xlsx workbook; never copied into the repository.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Sanitized template JSON destination.")
    args = parser.parse_args()
    try:
        template = extract_template(args.workbook)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(template, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    except (OSError, ValueError) as error:
        parser.error(str(error))
    print(
        f"Extracted {len(template['measurements'])} POMs; "
        f"{sum(row['required'] for row in template['measurements'])} required; "
        f"default unit {template['defaultUnit']}; "
        f"{template['source']['formulaCount']} formulas."
    )
    print("Source palette:", json.dumps(template["palette"], sort_keys=True))
    print("Sanitized template:", args.output)


if __name__ == "__main__":
    main()
